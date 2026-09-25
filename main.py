from datetime import datetime
import io
import json
import os
import smtplib
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
import warnings
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd
import requests
import yfinance as yf

warnings.filterwarnings('ignore')

# =====================================================================
# CONFIGURAZIONE E PARAMETRI MASTER
# =====================================================================
SENDER_EMAIL = os.environ.get('SENDER_EMAIL')
APP_PASSWORD = os.environ.get('APP_PASSWORD')
RECEIVER_EMAIL = os.environ.get('RECEIVER_EMAIL', SENDER_EMAIL)
SOGLIA_STORNO_MINIMA = (
    15.0  # Filtro base: il titolo deve essere in sconto di almeno il 15%
)
STATE_FILE = 'master_screener_stato_precedente.json'


# =====================================================================
# 0. MEMORIA STORICA CAMBIAMENTI
# =====================================================================
def carica_stato_precedente():
  if os.path.exists(STATE_FILE):
    try:
      with open(STATE_FILE, 'r', encoding='utf-8') as f:
        return json.load(f)
    except Exception:
      return {}
  return {}


def salva_stato_attuale(df_res):
  nuovo_stato = {}
  for _, row in df_res.iterrows():
    ticker = row['Ticker']
    top10 = row['TOP 10 OCCASIONI']
    verdetto = row['VERDETTO DEL CONSULENTE (AZIONE RAPIDA)']
    nuovo_stato[ticker] = {'top10': top10, 'verdetto': verdetto}
  with open(STATE_FILE, 'w', encoding='utf-8') as f:
    json.dump(nuovo_stato, f, indent=4, ensure_ascii=False)


# =====================================================================
# 1. RECUPERO TICKER (S&P 500 + NASDAQ 100 + MIDCAP 400)
# =====================================================================
def _fetch_wikipedia_table(url):
  headers = {
      'User-Agent': (
          'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
      )
  }
  response = requests.get(url, headers=headers, timeout=15)
  return pd.read_html(io.StringIO(response.text))


def ottieni_ticker_usa():
  tickers = set()
  try:
    df_sp = pd.read_csv(
        'https://raw.githubusercontent.com/datasets/s-and-p-500-companies/main/data/constituents.csv'
    )
    tickers.update(
        df_sp['Symbol']
        .str.replace('.', '-', regex=False)
        .str.strip()
        .tolist()
    )
  except Exception as e:
    print(f'Errore recupero S&P 500: {e}')
  try:
    tables = _fetch_wikipedia_table('https://en.wikipedia.org/wiki/Nasdaq-100')
    for df in tables:
      col = next(
          (
              c
              for c in df.columns
              if str(c).lower()
              in ['ticker', 'symbol', 'company stock symbol']
          ),
          None,
      )
      if col:
        tickers.update(
            df[col]
            .dropna()
            .astype(str)
            .str.replace('.', '-', regex=False)
            .str.strip()
            .tolist()
        )
        break
  except Exception as e:
    print(f'Errore recupero Nasdaq 100: {e}')
  try:
    tables = _fetch_wikipedia_table(
        'https://en.wikipedia.org/wiki/List_of_S%26P_400_companies'
    )
    for df in tables:
      col = next(
          (
              c
              for c in df.columns
              if str(c).lower() in ['symbol', 'ticker', 'company stock symbol']
          ),
          None,
      )
      if col:
        tickers.update(
            df[col]
            .dropna()
            .astype(str)
            .str.replace('.', '-', regex=False)
            .str.strip()
            .tolist()
        )
        break
  except Exception as e:
    print(f'Errore recupero S&P MidCap 400: {e}')
  return list(tickers)


# =====================================================================
# 2. MOTORE DI CALCOLO INDICATORI AVANZATI
# =====================================================================
def calcola_rsi(series, period=14):
  delta = series.diff()
  gain = (delta.where(delta > 0, 0)).rolling(window=period).mean()
  loss = (-delta.where(delta < 0, 0)).rolling(window=period).mean()
  rs = gain / (loss + 1e-10)
  return 100 - (100 / (1 + rs))


def calcola_obv(chiusure, volumi):
  return (np.sign(chiusure.diff()) * volumi).fillna(0).cumsum()


def calcola_cmf(massimi, minimi, chiusure, volumi, periodi=20):
  mf_multiplier = ((chiusure - minimi) - (massimi - chiusure)) / (
      massimi - minimi + 1e-10
  )
  mf_volume = mf_multiplier * volumi
  return mf_volume.rolling(window=periodi).sum() / (
      volumi.rolling(window=periodi).sum() + 1e-10
  )


def calcola_close_location_value(chiusura, minimo, massimo):
  rng = massimo - minimo
  if rng == 0:
    return 0.5
  return (chiusura - minimo) / rng


def calcola_volume_poc(chiusure, volumi, periodi=60, bins=10):
  if len(chiusure) < periodi:
    return float(chiusure.iloc[-1])
  sub_close = chiusure.iloc[-periodi:]
  sub_vol = volumi.iloc[-periodi:]
  counts, bin_edges = np.histogram(sub_close, bins=bins, weights=sub_vol)
  max_idx = np.argmax(counts)
  return float((bin_edges[max_idx] + bin_edges[max_idx + 1]) / 2.0)


def rileva_divergenza_cmf(df, lookback=15):
  if len(df) < lookback + 5:
    return 'Assente'
  sub_df = df.iloc[-lookback:].copy()
  prices = sub_df['Close'].values
  cmf_vals = sub_df['CMF'].values
  p_min1, p_min2 = np.min(prices[: lookback // 2]), np.min(
      prices[lookback // 2 :]
  )
  c_min1, c_min2 = np.min(cmf_vals[: lookback // 2]), np.min(
      cmf_vals[lookback // 2 :]
  )
  p_max1, p_max2 = np.max(prices[: lookback // 2]), np.max(
      prices[lookback // 2 :]
  )
  c_max1, c_max2 = np.max(cmf_vals[: lookback // 2]), np.max(
      cmf_vals[lookback // 2 :]
  )
  if p_min2 < p_min1 and c_min2 > c_min1:
    return 'Rialzista 🟢'
  if p_max2 > p_max1 and c_max2 < c_max1:
    return 'Ribassista 🔴'
  return 'Assente'


def calcola_volatilita_squeeze(df, length=20):
  sma = df['Close'].rolling(window=length).mean()
  std = df['Close'].rolling(window=length).std()
  bb_upper, bb_lower = sma + (2 * std), sma - (2 * std)
  atr = (
      (df['High'] - df['Low'])
      .combine((df['High'] - df['Close'].shift()).abs(), max)
      .combine((df['Low'] - df['Close'].shift()).abs(), max)
      .rolling(window=length)
      .mean()
  )
  kc_upper, kc_lower = sma + (1.5 * atr), sma - (1.5 * atr)
  return (bb_upper.iloc[-1] <= kc_upper.iloc[-1]) and (
      bb_lower.iloc[-1] >= kc_lower.iloc[-1]
  )


# =====================================================================
# 3. FILTRO FONDAMENTALE INTELLIGENTE
# =====================================================================
def verifica_salute_finanziaria_intelligente(t_obj, cmf_val):
  try:
    info = t_obj.info or {}
    revenue_growth = info.get('revenueGrowth', None)
    earnings_growth = info.get('earningsGrowth', None)
    debt_to_equity = info.get('debtToEquity', None)

    if revenue_growth is not None and revenue_growth < -0.30 and cmf_val < 0.10:
      return False, 'Crollo pesante dei ricavi senza supporto istituzionale.'

    if debt_to_equity is not None and debt_to_equity > 400 and cmf_val < 0.05:
      return False, 'Indice di indebitamento critico senza flussi a favore.'

    if earnings_growth is not None and earnings_growth < -0.35 and cmf_val < 0.05:
      return False, 'Utili in forte calo senza interesse delle mani forti.'

    return True, 'Azienda solida o in fase di investimento strategico.'
  except Exception:
    return True, 'Dati fondamentali parziali, validato da analisi tecnica.'


# =====================================================================
# 3.5 NUOVA FUNZIONE: VALUTAZIONE IDONEITÀ PER PORTAFOGLIO PERSONALE
# =====================================================================
def valuta_idoneita_portafoglio(
    info, cmf_val, vsa_rating, storno_pct, is_bear_market
):
  """Valuta se il titolo merita di passare nel Portafoglio Personale per un accumulo PAC / Lungo Termine."""
  try:
    fwd_pe = info.get('forwardPE', None)
    peg = info.get('pegRatio', None)
    debt_to_equity = info.get('debtToEquity', None)
    short_pct = (
        info.get('shortPercentOfFloat', 0) * 100
        if info.get('shortPercentOfFloat')
        else 0
    )

    # Condizione 1: Eccessivamente speculativo (es. Short Squeeze estremo senza fondamentali sani)
    if short_pct > 12 and (fwd_pe is None or fwd_pe > 40):
      return '🔴 SOLO TRADE BREVE: Volatilità da squeeze speculativo. Non idoneo al cassetto.'

    # Condizione 2: Ottima azienda a sconto con flussi di denaro veri
    punti_qualita = 0
    if fwd_pe is not None and 0 < fwd_pe < 25:
      punti_qualita += 1
    if peg is not None and 0 < peg < 1.5:
      punti_qualita += 1
    if cmf_val > 0.05 or vsa_rating == '🟢 ACCUMULAZIONE PULITA':
      punti_qualita += 1
    if storno_pct >= 20.0:
      punti_qualita += 1

    if punti_qualita >= 3:
      return '🟢 IDONEO: Azienda solida a sconto, ottima candidatura da passare in Portafoglio Personale!'
    elif punti_qualita == 2:
      return '🟡 CON RISERVA: Buon rimbalzo di breve. Verifica i dati di bilancio prima di un PAC.'
    else:
      return (
          '🔴 SOLO TRADE BREVE: Buona dinamica tecnica/volumi, ma poco adatta'
          ' al lungo termine.'
      )
  except Exception:
    return (
        '🟡 NEUTRO: Dati insufficienti per valutare il lungo termine, gestisci'
        ' come trade breve.'
    )


# =====================================================================
# 4. ANALISI DEL SINGOLO TITOLO & COSTRUZIONE VERDETTO
# =====================================================================
def analizza_titolo(ticker_str, data_oggi):
  try:
    t = yf.Ticker(ticker_str)
    df = t.history(period='1y')
    if df.empty or len(df) < 200:
      return None
    if isinstance(df.columns, pd.MultiIndex):
      df.columns = df.columns.get_level_values(0)

    # Indicatori
    df['CMF'] = calcola_cmf(df['High'], df['Low'], df['Close'], df['Volume'])
    df['RSI'] = calcola_rsi(df['Close'])
    df['EMA20'] = df['Close'].ewm(span=20, adjust=False).mean()
    df['SMA50'] = df['Close'].rolling(window=50).mean()
    df['SMA200'] = df['Close'].rolling(window=200).mean()

    curr = df.iloc[-1]
    prezzo_attuale = float(curr['Close'])
    massimo_52w = float(df['High'].max())
    minimo_52w = float(df['Low'].min())
    supporto_60g = float(df['Close'].iloc[-60:].min())
    minimo_60g = supporto_60g

    # 1. Filtro storno minimo 15% dai massimi
    storno_pct = ((massimo_52w - prezzo_attuale) / massimo_52w) * 100
    if storno_pct < SOGLIA_STORNO_MINIMA:
      return None

    # 2. Salute finanziaria flessibile
    cmf_corrente = float(curr['CMF'])
    is_sano, nota_bilancio = verifica_salute_finanziaria_intelligente(
        t, cmf_corrente
    )
    if not is_sano:
      return None

    # Info anagrafiche e fondamentali
    info = t.info or {}
    nome_azienda = info.get('shortName') or info.get('longName') or ''
    ticker_display = (
        f'{ticker_str} - {nome_azienda}' if nome_azienda else ticker_str
    )
    target_price = info.get('targetMeanPrice', None)
    upside = (
        round(((target_price - prezzo_attuale) / prezzo_attuale) * 100, 2)
        if target_price
        else 'N/D'
    )
    fwd_pe = round(info.get('forwardPE'), 2) if info.get('forwardPE') else 'N/D'
    peg = round(info.get('pegRatio'), 2) if info.get('pegRatio') else 'N/D'
    short_pct = (
        round(info.get('shortPercentOfFloat', 0) * 100, 2)
        if info.get('shortPercentOfFloat')
        else 'N/D'
    )

    sma_50 = float(curr['SMA50'])
    sma_200 = float(curr['SMA200'])
    dist_sma50_pct = round(((prezzo_attuale - sma_50) / sma_50) * 100, 2)
    dist_sma200_pct = round(((prezzo_attuale - sma_200) / sma_200) * 100, 2)
    distanza_minimo_60g_pct = round(
        ((prezzo_attuale - minimo_60g) / minimo_60g) * 100, 2
    )
    is_bear_market = prezzo_attuale < sma_200
    divergenza = rileva_divergenza_cmf(df)
    is_squeeze = calcola_volatilita_squeeze(df)
    media_vol_5g = df['Volume'].iloc[-5:].mean()
    media_vol_60g = df['Volume'].iloc[-60:].mean()
    rvol_5d_pct = (
        round((media_vol_5g / media_vol_60g) * 100, 1)
        if media_vol_60g > 0
        else 0.0
    )

    obv_serie = calcola_obv(df['Close'], df['Volume'])
    obv_sma = obv_serie.rolling(20).mean()
    obv_trend = (
        'Rialzista (Accumulo)'
        if float(obv_serie.iloc[-1]) > float(obv_sma.iloc[-1])
        else 'Ribassista (Distribuzione)'
    )

    clv_5d = [
        calcola_close_location_value(
            df['Close'].iloc[i], df['Low'].iloc[i], df['High'].iloc[i]
        )
        for i in range(-5, 0)
    ]
    clv_val = float(np.mean(clv_5d))
    poc_val = calcola_volume_poc(df['Close'], df['Volume'], periodi=60)

    if cmf_corrente > 0.05 and clv_val >= 0.55:
      vsa_rating = '🟢 ACCUMULAZIONE PULITA'
    elif cmf_corrente < -0.05 and clv_val <= 0.45:
      vsa_rating = '🔴 DISTRIBUZIONE'
    else:
      vsa_rating = '🟡 NEUTRO'

    # --- VERDETTI DIRETTI OPERATIVI (BREVE TERMINE/EXPLOSIVE) ---
    if (
        short_pct != 'N/D'
        and float(short_pct) > 10
        and cmf_corrente > 0.10
        and is_squeeze
    ):
      verdetto = (
          '🔥 [BREVE] OCCASIONE SHORT SQUEEZE: Esplosivo, volumi alle stelle!'
          ' Pronto al fuoco.'
      )
      orizzonte = 'Breve Termine (Esplosivo)'
    elif divergenza == 'Rialzista 🟢':
      verdetto = (
          "🚀 [BREVE/MEDIO] OCCASIONE D'ORO: Divergenza rialzista sui minimi,"
          ' timing perfetto per un possibile trend di brevissimo!'
      )
      orizzonte = 'Breve/Medio Termine'
    elif distanza_minimo_60g_pct <= 3.0 and cmf_corrente >= 0.0:
      verdetto = (
          '💎 [ACCUMULO] VICINO AI MINIMI: Entra perché quest\'azienda è sana e'
          ' non ti ricapita a questi prezzi.'
      )
      orizzonte = 'Lungo Termine (Accumulo Silenzioso)'
    elif not is_bear_market and cmf_corrente > 0.05:
      verdetto = (
          '🛡️ [CASSETTO] CARRO ARMATO IN BULL TREND: Struttura solida, cammina'
          ' piano ma viaggia sicura sul lungo.'
      )
      orizzonte = 'Lungo Termine (Cassetto)'
    elif is_bear_market and cmf_corrente > 0.05:
      verdetto = (
          '💎 [PAC] SCONTO PROFONDO: Sotto SMA200 ma le mani forti stanno'
          ' accumulando sul ribasso. Ottimo affare.'
      )
      orizzonte = 'Lungo Termine (PAC a Sconto)'
    else:
      verdetto = (
          f'🔍 [ATTESA] MONITORARE: Struttura incerta ({vsa_rating}). Nota:'
          f' {nota_bilancio}'
      )
      orizzonte = 'Monitoraggio / Attendere'

    # --- SUGGERIMENTO PER PORTAFOGLIO PERSONALE ---
    suggerimento_portafoglio = valuta_idoneita_portafoglio(
        info, cmf_corrente, vsa_rating, storno_pct, is_bear_market
    )

    # Score interno per la classifica (Mantiene in alto i titoli con exploit di breve imminente)
    score = 50
    if cmf_corrente > 0.10:
      score += 20
    elif cmf_corrente > 0:
      score += 10
    if curr['Close'] > curr['EMA20']:
      score += 10
    if vsa_rating == '🟢 ACCUMULAZIONE PULITA':
      score += 15
    if divergenza == 'Rialzista 🟢':
      score += 15
    score = max(0, min(100, int(score)))

    return {
        'TOP 10 OCCASIONI': '-',
        'Ticker': ticker_display,
        'VERDETTO DEL CONSULENTE (AZIONE RAPIDA)': verdetto,
        'Idoneità Portafoglio Personale (PAC/Lungo)': suggerimento_portafoglio,
        'Orizzonte Strategico': orizzonte,
        'Prezzo Attuale ($)': round(prezzo_attuale, 2),
        'Supporto 60G ($)': round(supporto_60g, 2),
        'Distanza dal Minimo 60G (%)': distanza_minimo_60g_pct,
        'Resistenza Max 52W ($)': round(massimo_52w, 2),
        'Storno dai Max 52W (%)': round(storno_pct, 1),
        'Distanza da SMA 50 (%)': dist_sma50_pct,
        'Distanza da SMA 200 (%)': dist_sma200_pct,
        'Trend di Fondo': (
            '🔴 BEAR (Sotto SMA200)' if is_bear_market else '🟢 BULL (Sopra SMA200)'
        ),
        'Target Price Medio ($)': (
            round(target_price, 2) if target_price else 'N/D'
        ),
        'Upside Atteso (%)': upside,
        'RSI (14)': round(float(curr['RSI']), 1),
        'Chaikin Money Flow (CMF)': round(cmf_corrente, 3),
        'Divergenza CMF': divergenza,
        'Volumi 1W (% vs 60G)': rvol_5d_pct,
        'OBV Trend': obv_trend,
        'Analisi VSA (Flussi)': vsa_rating,
        'Squeeze Volatilità': 'Attivo 🔥' if is_squeeze else 'No',
        'Point of Control (POC 60G)': round(poc_val, 2),
        'Short Interest (%)': short_pct,
        'Forward P/E': fwd_pe,
        'PEG Ratio': peg,
        'Analisi Fondamentale / Nota': nota_bilancio,
        '_score_interno': score,
    }
  except Exception:
    return None


# =====================================================================
# 5. GENERAZIONE EXCEL PULITO E ORDINATO
# =====================================================================
def genera_excel(df_risultati):
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = 'Consulente Smart Money'
  cols_to_export = [c for c in df_risultati.columns if not c.startswith('_')]
  df_export = df_risultati[cols_to_export]

  for r in dataframe_to_rows(df_export, index=False, header=True):
    ws.append(r)

  header_fill = PatternFill(
      start_color='1F4E79', end_color='1F4E79', fill_type='solid'
  )
  top_fill = PatternFill(
      start_color='D9EAD3', end_color='D9EAD3', fill_type='solid'
  )
  header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
  thin_border = Border(
      left=openpyxl.styles.Side(style='thin', color='D9D9D9'),
      right=openpyxl.styles.Side(style='thin', color='D9D9D9'),
      top=openpyxl.styles.Side(style='thin', color='D9D9D9'),
      bottom=openpyxl.styles.Side(style='thin', color='D9D9D9'),
  )

  for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal='center', vertical='center', wrap_text=True
    )

  for row_idx, row in enumerate(
      ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column),
      start=2,
  ):
    is_top10 = row_idx <= 11
    for cell in row:
      cell.border = thin_border
      cell.alignment = Alignment(horizontal='center', vertical='center')
      if is_top10:
        cell.fill = top_fill

  for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

  data_oggi = datetime.now().strftime('%Y-%m-%d')
  excel_file = f'Consulente_Smart_Money_{data_oggi}.xlsx'
  wb.save(excel_file)
  return excel_file


# =====================================================================
# 6. INVIO E-MAIL CON REPORT EXCEL
# =====================================================================
def invia_email_report(file_path, data_oggi, num_titoli):
  if not all([SENDER_EMAIL, APP_PASSWORD]):
    print(
        '⚠️ Credenziali e-mail non configurate. File Excel salvato nella'
        ' cartella locale.'
    )
    return
  msg = MIMEMultipart()
  msg['From'] = SENDER_EMAIL
  msg['To'] = RECEIVER_EMAIL
  msg['Subject'] = (
      f'🧠 Report Consulente Smart Money ({data_oggi}) - Trovate {num_titoli}'
      ' Occasioni'
  )
  body = (
      f"Ciao! Il tuo consulente virtuale ha completato l'analisi in data"
      f' {data_oggi}.\n\nSono state filtrate le migliori occasioni di mercato'
      ' ordinate per rimbalzo/exploit di breve termine.\nÈ stata integrata la'
      ' colonna per identificare i titoli idonei al tuo Portafoglio'
      ' Personale.\n\nBuon gain!'
  )
  msg.attach(MIMEText(body, 'plain', 'utf-8'))
  with open(file_path, 'rb') as f:
    part = MIMEApplication(f.read(), Name=os.path.basename(file_path))
    part['Content-Disposition'] = (
        f'attachment; filename="{os.path.basename(file_path)}"'
    )
    msg.attach(part)
  try:
    server = smtplib.SMTP('smtp.gmail.com', 587)
    server.starttls()
    server.login(SENDER_EMAIL, APP_PASSWORD)
    server.send_message(msg)
    server.quit()
    print('✅ E-mail con report Excel inviata con successo!')
  except Exception as e:
    print(f'❌ Errore invio e-mail: {e}')


# =====================================================================
# 7. ESECUZIONE PRINCIPALE
# =====================================================================
if __name__ == '__main__':
  data_oggi = datetime.now().strftime('%Y-%m-%d')
  print(f'=== Avvio Consulente Smart Money & Screener ({data_oggi}) ===')
  tickers = ottieni_ticker_usa()
  print(f'Titoli totali in scansione: {len(tickers)}')
  risultati = []
  for idx, t in enumerate(tickers):
    res = analizza_titolo(t, data_oggi)
    if res:
      risultati.append(res)
    if (idx + 1) % 100 == 0:
      print(f'Analizzati {idx + 1}/{len(tickers)}...')
  if risultati:
    df_res = pd.DataFrame(risultati)
    df_res = df_res.sort_values(
        by=['_score_interno', 'Chaikin Money Flow (CMF)'], ascending=False
    ).reset_index(drop=True)
    for i in range(min(10, len(df_res))):
      df_res.at[i, 'TOP 10 OCCASIONI'] = f'⭐ TOP {i+1}'
    df_res = df_res.drop(columns=['_score_interno'])
    excel_path = genera_excel(df_res)
    invia_email_report(excel_path, data_oggi, len(df_res))
    print(
        f'🚀 Analisi completata! Il consulente ha selezionato {len(df_res)}'
        ' occasioni salvate nel file Excel.'
    )
  else:
    print('Nessun titolo rispetta i criteri del consulente per oggi.')

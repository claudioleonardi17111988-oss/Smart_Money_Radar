#!/usr/bin/env python3
"""
MONEY FLOW PRO v2 - Screener / consulente su S&P 500 + Nasdaq 100 + S&P MidCap 400.

Pensato per girare gratis su GitHub Actions usando solo yfinance (nessun abbonamento).

Pipeline:
  1. Universo ticker (con cache "ultima lista valida" se le fonti falliscono)
  2. Prezzi 1y in batch (yf.download) + validazione (barra parziale, dati stale, liquidita')
  3. Filtro tecnico: storno dai massimi 52w compreso tra [storno_min, storno_max]
  4. Fondamentali SOLO per i candidati, con cache su disco (TTL 7gg), retry e circuit breaker
  5. Filtro qualita' (anti value-trap) + Fair Value multi-modello con indice di confidenza
  6. Verdetto per regole + score a percentili cross-sectional
  7. Top N (solo verdetti azionabili, cap per settore), Excel formattato, email HTML
  8. Storico giornaliero su CSV per il forward test (vedi forward_test.py)

Uso locale:   python money_flow_pro.py --limit 60 --no-email
Variabili d'ambiente (GitHub Secrets): SENDER_EMAIL, APP_PASSWORD, RECEIVER_EMAIL (opzionale)

Nota: strumento di ricerca/screening a scopo informativo, non e' consulenza finanziaria.
"""
from __future__ import annotations

import argparse
import html
import io
import json
import logging
import math
import os
import random
import re
import smtplib
import ssl
import sys
import threading
import time
import traceback
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import datetime, time as dtime, timedelta, timezone
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from enum import Enum
from pathlib import Path
from typing import Optional
from zoneinfo import ZoneInfo

import numpy as np
import openpyxl
import pandas as pd
import requests
import yfinance as yf
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger("money_flow")

# =====================================================================
# CONFIGURAZIONE
# =====================================================================
SENDER_EMAIL = os.environ.get("SENDER_EMAIL")
APP_PASSWORD = os.environ.get("APP_PASSWORD")
RECEIVER_EMAIL = os.environ.get("RECEIVER_EMAIL") or SENDER_EMAIL

DATA_DIR = Path(os.environ.get("DATA_DIR", "data"))
CACHE_DIR = DATA_DIR / "cache" / "fund"          # cache fondamentali (NON committata)
UNIVERSE_FILE = DATA_DIR / "universe_cache.json"   # ultima lista ticker valida
HISTORY_FILE = DATA_DIR / "picks_history.csv"      # storico per forward test (committato)
NY = ZoneInfo("America/New_York")


@dataclass(frozen=True)
class Config:
    # --- universo / qualita' dati
    min_universe: int = 400            # sotto questa soglia la lista ticker e' considerata rotta
    min_bars: int = 200
    min_price: float = 5.0
    min_dollar_volume: float = 5e6     # media 20g di Close*Volume
    stale_days: int = 4                # ticker con ultima barra piu' vecchia di SPY oltre N giorni = scartato
    # --- filtro tecnico
    storno_min: float = 15.0           # % sotto i max 52w
    storno_max: float = 70.0           # oltre: probabile titolo in dissesto
    # --- download
    batch_size: int = 80
    batch_pause: float = 1.5
    info_workers: int = 4
    fund_ttl_days: int = 7
    fund_max_consecutive_fail: int = 15   # circuit breaker (Yahoo che throttla)
    # --- indicatori
    vol_slope_thr: float = 0.5         # %/giorno sul log-volume (Theil-Sen)
    # --- fair value
    fcf_required_yield: float = 0.06
    fv_min_ratio: float = 0.25         # un modello con FV < 25% o > 400% del prezzo e' scartato come outlier
    fv_max_ratio: float = 4.0
    min_analysts: int = 5
    # --- output
    top_n: int = 10
    max_per_sector: int = 3
    only_actionable_top: bool = True
    earnings_warn_days: int = 7


CFG = Config()

# =====================================================================
# ENUM (la logica NON dipende piu' da stringhe con emoji)
# =====================================================================


class Vsa(str, Enum):
    ACCUMULO = "🟢 ACCUMULAZIONE PULITA"
    DISTRIB = "🔴 DISTRIBUZIONE"
    NEUTRO = "🟡 NEUTRO"


class Div(str, Enum):
    RIALZ = "Rialzista 🟢"
    RIBASS = "Ribassista 🔴"
    NO = "Assente"


class Tr5(str, Enum):
    ACCEL = "In Accelerazione 📈"
    RAFF = "In Raffreddamento 📉"
    STAB = "Stabile ➡️"


class VolReg(str, Enum):
    CRESC = "In Crescita 📈"
    ESAUR = "In Esaurimento 📉"
    STAB = "Stabile ➡️"
    ND = "N/D"


class Cat(str, Enum):
    BREVE_SQUEEZE = "BREVE (squeeze)"
    BREVE_MEDIO = "BREVE/MEDIO"
    ACCUMULO = "ACCUMULO"
    CASSETTO = "CASSETTO"
    PAC = "PAC"
    ATTESA = "ATTESA"
    EVITARE = "EVITARE"


AZIONABILI = {Cat.BREVE_SQUEEZE, Cat.BREVE_MEDIO, Cat.ACCUMULO, Cat.CASSETTO, Cat.PAC}
ORIZZONTE = {
    Cat.BREVE_SQUEEZE: "Breve Termine (Speculativo)",
    Cat.BREVE_MEDIO: "Breve/Medio Termine",
    Cat.ACCUMULO: "Lungo Termine (Accumulo Silenzioso)",
    Cat.CASSETTO: "Lungo Termine (Cassetto)",
    Cat.PAC: "Lungo Termine (PAC a Sconto)",
    Cat.ATTESA: "Monitoraggio / Attendere",
    Cat.EVITARE: "Evitare per ora",
}

SETTORI_IT = {
    "Technology": "Tecnologia dell'Informazione",
    "Industrials": "Industriale",
    "Financial Services": "Finanziari",
    "Consumer Cyclical": "Beni di Consumo Discrezionali",
    "Healthcare": "Salute",
    "Utilities": "Servizi di Pubblica Utilità",
    "Energy": "Energia",
    "Consumer Defensive": "Beni di Consumo Primari",
    "Communication Services": "Servizi di Comunicazione",
    "Real Estate": "Real Estate (REITs)",
    "Basic Materials": "Materiali di Base",
}

# P/E "prudenziali" per settore (euristiche modificabili, NON verita' di mercato)
PE_BASE = {
    "Technology": 22, "Healthcare": 18, "Consumer Cyclical": 18, "Communication Services": 17,
    "Industrials": 17, "Financial Services": 12, "Utilities": 16, "Energy": 11,
    "Consumer Defensive": 18, "Basic Materials": 13,
}
SETTORI_GROWTH = {"Technology", "Healthcare", "Consumer Cyclical", "Communication Services"}
SETTORI_CICLICI = {"Energy", "Basic Materials"}

FUND_KEYS = [
    "sector", "industry", "shortName", "longName", "forwardEps", "trailingEps", "earningsGrowth",
    "revenueGrowth", "debtToEquity", "bookValue", "returnOnEquity", "operatingMargins",
    "profitMargins", "freeCashflow", "totalCash", "totalDebt", "totalRevenue", "ebitda",
    "forwardPE", "trailingPE", "pegRatio", "targetMeanPrice", "targetLowPrice", "targetHighPrice",
    "numberOfAnalystOpinions", "shortPercentOfFloat", "marketCap", "sharesOutstanding",
    "currentRatio", "earningsTimestamp", "earningsTimestampStart", "recommendationMean",
]


# =====================================================================
# UTILITA'
# =====================================================================
def num(x) -> Optional[float]:
    """Converte in float finito, altrimenti None (Yahoo restituisce None, stringhe, NaN, inf...)."""
    try:
        if x is None or isinstance(x, bool):
            return None
        v = float(x)
        return v if math.isfinite(v) else None
    except (TypeError, ValueError):
        return None


def clamp(x: float, lo: float, hi: float) -> float:
    return max(lo, min(hi, x))


def retry(fn, tries: int = 3, base: float = 2.0):
    """Retry con backoff esponenziale; attesa piu' lunga se sembra un rate limit di Yahoo."""
    for attempt in range(tries):
        try:
            return fn()
        except Exception as exc:  # noqa: BLE001 - vogliamo riprovare qualunque errore di rete/parsing
            if attempt == tries - 1:
                raise
            msg = f"{type(exc).__name__} {exc}".lower()
            rate = any(k in msg for k in ("rate", "429", "too many"))
            wait = base * (2 ** attempt) * (3 if rate else 1) + random.random()
            log.debug("retry %d/%d tra %.1fs (%s)", attempt + 1, tries, wait, msg[:80])
            time.sleep(wait)


# =====================================================================
# 1. UNIVERSO TICKER
# =====================================================================
def _wiki_tables(url: str):
    r = requests.get(url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}, timeout=20)
    r.raise_for_status()
    return pd.read_html(io.StringIO(r.text))


def _colonna_ticker(tables, nomi):
    for df in tables:
        col = next((c for c in df.columns if str(c).strip().lower() in nomi), None)
        if col is not None:
            return df[col]
    return None


def _pulisci_simboli(serie) -> set:
    out = set()
    for s in serie.dropna().astype(str):
        s = s.strip().upper().replace(".", "-")
        if re.fullmatch(r"[A-Z0-9\-]{1,8}", s):
            out.add(s)
    return out


def ottieni_universo() -> list[str]:
    tickers: set[str] = set()
    fonti = {
        "S&P 500": lambda: pd.read_csv(
            "https://raw.githubusercontent.com/datasets/s-and-p-500-companies/main/data/constituents.csv"
        )["Symbol"],
        "Nasdaq 100": lambda: _colonna_ticker(
            _wiki_tables("https://en.wikipedia.org/wiki/Nasdaq-100"), {"ticker", "symbol"}),
        "S&P MidCap 400": lambda: _colonna_ticker(
            _wiki_tables("https://en.wikipedia.org/wiki/List_of_S%26P_400_companies"), {"symbol", "ticker"}),
    }
    for nome, fn in fonti.items():
        try:
            serie = retry(fn, tries=3, base=2)
            if serie is None:
                raise ValueError("colonna ticker non trovata")
            s = _pulisci_simboli(serie)
            log.info("Universo %-15s %d simboli", nome, len(s))
            tickers |= s
        except Exception as exc:  # noqa: BLE001
            log.warning("Fonte %s fallita: %s", nome, exc)

    if len(tickers) >= CFG.min_universe:
        UNIVERSE_FILE.parent.mkdir(parents=True, exist_ok=True)
        UNIVERSE_FILE.write_text(json.dumps(sorted(tickers)), encoding="utf-8")
        return sorted(tickers)

    log.error("Universo troppo piccolo (%d < %d): provo la cache.", len(tickers), CFG.min_universe)
    if UNIVERSE_FILE.exists():
        cached = json.loads(UNIVERSE_FILE.read_text(encoding="utf-8"))
        if len(cached) >= CFG.min_universe:
            return sorted(set(cached) | tickers)
    raise RuntimeError("Impossibile costruire l'universo dei ticker (fonti KO e nessuna cache valida).")


# =====================================================================
# 2. PREZZI: DOWNLOAD IN BATCH + VALIDAZIONE
# =====================================================================
def _estrai_batch(raw: pd.DataFrame, batch: list[str]) -> dict[str, pd.DataFrame]:
    res: dict[str, pd.DataFrame] = {}
    if raw is None or raw.empty:
        return res
    if isinstance(raw.columns, pd.MultiIndex):
        lvl0 = set(raw.columns.get_level_values(0))
        if "Close" in lvl0 and not any(t in lvl0 for t in batch):   # livelli invertiti
            raw = raw.swaplevel(0, 1, axis=1)
            lvl0 = set(raw.columns.get_level_values(0))
        for tk in batch:
            if tk in lvl0:
                d = raw[tk].dropna(how="all")
                if not d.empty:
                    res[tk] = d
    elif len(batch) == 1:
        res[batch[0]] = raw.dropna(how="all")
    return res


def scarica_prezzi(tickers: list[str], stats: Counter, batch_size: Optional[int] = None,
                   **dl_kwargs) -> dict[str, pd.DataFrame]:
    """Scarica i prezzi a blocchi, con retry; un secondo passaggio riprova i mancanti in blocchi piu' piccoli."""
    if "start" not in dl_kwargs:
        dl_kwargs.setdefault("period", "1y")
    bs = batch_size or CFG.batch_size
    out: dict[str, pd.DataFrame] = {}

    def _scarica(batch):
        def _call():
            raw = yf.download(batch, interval="1d", auto_adjust=True, group_by="ticker",
                              threads=True, progress=False, **dl_kwargs)
            if raw is None or raw.empty:
                raise ValueError("download vuoto")
            return raw
        return _estrai_batch(retry(_call, tries=3, base=3), batch)

    def _passaggio(lista, dim):
        for i in range(0, len(lista), dim):
            batch = lista[i:i + dim]
            try:
                out.update(_scarica(batch))
            except Exception as exc:  # noqa: BLE001
                stats["batch_prezzi_falliti"] += 1
                log.warning("Batch prezzi %d-%d fallito: %s", i, i + len(batch), exc)
            time.sleep(CFG.batch_pause)

    _passaggio(tickers, bs)
    mancanti = [t for t in tickers if t not in out]
    if mancanti:
        log.info("Secondo passaggio prezzi su %d ticker mancanti", len(mancanti))
        _passaggio(mancanti, max(10, bs // 3))
    stats["prezzi_mancanti"] += len([t for t in tickers if t not in out])
    return out


def pulisci_ohlcv(df: pd.DataFrame, ora_ny: datetime) -> Optional[pd.DataFrame]:
    """Normalizza colonne/indice e rimuove la barra odierna se il mercato non ha ancora chiuso."""
    cols = ["Open", "High", "Low", "Close", "Volume"]
    if df is None or any(c not in df.columns for c in cols):
        return None
    d = df[cols].copy()
    if getattr(d.index, "tz", None) is not None:
        d.index = d.index.tz_localize(None)
    d = d.dropna(subset=["High", "Low", "Close"])
    d["Volume"] = d["Volume"].fillna(0)
    if d.empty:
        return None
    mercato_aperto = ora_ny.weekday() < 5 and ora_ny.time() < dtime(16, 30)
    if d.index[-1].date() == ora_ny.date() and mercato_aperto:
        d = d.iloc[:-1]           # candela parziale: falserebbe CMF, CLV e volumi
    return d if not d.empty else None


# =====================================================================
# 3. INDICATORI
# =====================================================================
def rsi_wilder(close: pd.Series, period: int = 14) -> pd.Series:
    d = close.diff()
    gain = d.clip(lower=0).ewm(alpha=1 / period, adjust=False, min_periods=period).mean()
    loss = (-d.clip(upper=0)).ewm(alpha=1 / period, adjust=False, min_periods=period).mean()
    return 100 - 100 / (1 + gain / loss)


def clv_series(h, l, c) -> pd.Series:
    return ((c - l) / (h - l).replace(0, np.nan)).fillna(0.5)


def cmf_series(h, l, c, v, n: int = 20) -> pd.Series:
    mfm = ((2 * c - l - h) / (h - l).replace(0, np.nan)).fillna(0)
    return (mfm * v).rolling(n).sum() / v.rolling(n).sum().replace(0, np.nan)


def true_range(df: pd.DataFrame) -> pd.Series:
    pc = df["Close"].shift()
    return pd.concat([df["High"] - df["Low"], (df["High"] - pc).abs(), (df["Low"] - pc).abs()], axis=1).max(axis=1)


def squeeze_attivo(df: pd.DataFrame, n: int = 20) -> bool:
    c = df["Close"]
    sma, sd = c.rolling(n).mean(), c.rolling(n).std()
    atr = true_range(df).rolling(n).mean()
    return bool(
        (sma + 2 * sd).iloc[-1] < (sma + 1.5 * atr).iloc[-1]
        and (sma - 2 * sd).iloc[-1] > (sma - 1.5 * atr).iloc[-1]
    )


def theil_sen_pct(vol: pd.Series, n: int) -> float:
    """Pendenza robusta (Theil-Sen) del log-volume, in % al giorno. Insensibile ai singoli spike (earnings)."""
    y = np.log1p(vol.iloc[-n:].to_numpy(dtype=float))
    if len(y) < n or not np.isfinite(y).all():
        return float("nan")
    i, j = np.triu_indices(n, 1)
    return float(np.median((y[j] - y[i]) / (j - i)) * 100)


def classifica_reg(slope: float) -> VolReg:
    if not math.isfinite(slope):
        return VolReg.ND
    if slope > CFG.vol_slope_thr:
        return VolReg.CRESC
    if slope < -CFG.vol_slope_thr:
        return VolReg.ESAUR
    return VolReg.STAB


def _pivots(values: np.ndarray, k: int, kind: str) -> list[int]:
    idx = []
    for i in range(k, len(values) - k):
        w = values[i - k:i + k + 1]
        if not np.isfinite(values[i]):
            continue
        if (kind == "min" and values[i] == np.nanmin(w)) or (kind == "max" and values[i] == np.nanmax(w)):
            idx.append(i)
    return idx


def divergenza_cmf(df: pd.DataFrame, lookback: int = 45, k: int = 3, margin: float = 0.02) -> Div:
    """Divergenza prezzo/CMF su PIVOT veri (non min/max di meta' finestra)."""
    sub = df.iloc[-lookback:]
    if len(sub) < lookback:
        return Div.NO
    price, cm = sub["Close"].to_numpy(), sub["CMF"].to_numpy()
    for kind in ("min", "max"):
        idx = _pivots(price, k, kind)
        if len(idx) < 2:
            continue
        p2 = idx[-1]
        if p2 < lookback - 15:                      # il secondo pivot deve essere recente
            continue
        p1 = next((i for i in reversed(idx[:-1]) if p2 - i >= 5), None)
        if p1 is None:
            continue
        if kind == "min" and price[p2] < price[p1] and cm[p2] > cm[p1] + margin:
            return Div.RIALZ
        if kind == "max" and price[p2] > price[p1] and cm[p2] < cm[p1] - margin:
            return Div.RIBASS
    return Div.NO


def volume_poc(df: pd.DataFrame, n: int = 60, bins: int = 12) -> float:
    sub = df.iloc[-n:]
    tp = (sub["High"] + sub["Low"] + sub["Close"]) / 3
    counts, edges = np.histogram(tp, bins=bins, weights=sub["Volume"])
    i = int(np.argmax(counts))
    return float((edges[i] + edges[i + 1]) / 2)


def calcola_tecnici(df: pd.DataFrame, spy_ret20: float) -> Optional[dict]:
    c, h, l, v = df["Close"], df["High"], df["Low"], df["Volume"]
    d = df.copy()
    d["CMF"] = cmf_series(h, l, c, v)
    rsi = rsi_wilder(c)
    sma50, sma200 = c.rolling(50).mean(), c.rolling(200).mean()
    atr14 = true_range(df).ewm(alpha=1 / 14, adjust=False).mean()

    price = float(c.iloc[-1])
    cmf = float(d["CMF"].iloc[-1])
    chiavi = [cmf, float(rsi.iloc[-1]), float(sma50.iloc[-1]), float(sma200.iloc[-1]), float(atr14.iloc[-1])]
    if not all(math.isfinite(x) for x in chiavi):
        return None

    max52, min52 = float(h.max()), float(l.min())
    support60 = float(l.iloc[-60:].min())
    clv5 = float(clv_series(h, l, c).iloc[-5:].mean())

    if cmf > 0.05 and clv5 >= 0.55:
        vsa = Vsa.ACCUMULO
    elif cmf < -0.05 and clv5 <= 0.45:
        vsa = Vsa.DISTRIB
    else:
        vsa = Vsa.NEUTRO

    obv = (np.sign(c.diff()) * v).fillna(0).cumsum()
    obv_up = bool(obv.iloc[-1] > obv.rolling(20).mean().iloc[-1])

    dc, vv = c.diff().iloc[-20:], v.iloc[-20:]
    up_v, dn_v = float(vv[dc > 0].sum()), float(vv[dc < 0].sum())
    updown = clamp(up_v / dn_v, 0.2, 5.0) if dn_v > 0 else 5.0

    base_vol = float(v.iloc[-65:-5].median())
    rvol5 = float(v.iloc[-5:].mean() / base_vol) if base_vol > 0 else float("nan")
    if not math.isfinite(rvol5):
        trend5 = Tr5.STAB
    elif rvol5 >= 1.2:
        trend5 = Tr5.ACCEL
    elif rvol5 <= 0.8:
        trend5 = Tr5.RAFF
    else:
        trend5 = Tr5.STAB

    ret20 = float(c.iloc[-1] / c.iloc[-21] - 1)
    rsi_now = float(rsi.iloc[-1])
    return {
        "price": price, "max52": max52, "min52": min52,
        "storno": (max52 - price) / max52 * 100,
        "support60": support60,
        "dist_min60": (price / support60 - 1) * 100,
        "sma50": float(sma50.iloc[-1]), "sma200": float(sma200.iloc[-1]),
        "dist_sma50": (price / float(sma50.iloc[-1]) - 1) * 100,
        "dist_sma200": (price / float(sma200.iloc[-1]) - 1) * 100,
        "bear": price < float(sma200.iloc[-1]),
        "rsi": rsi_now,
        "rsi_rimbalzo": bool(rsi_now < 50 and rsi_now > float(rsi.iloc[-6]) + 3 and float(rsi.iloc[-15:].min()) < 35),
        "cmf": cmf, "clv5": clv5, "vsa": vsa, "obv_up": obv_up, "updown": updown,
        "rvol5": rvol5, "trend5": trend5,
        "reg20": classifica_reg(theil_sen_pct(v, 20)),
        "reg50": classifica_reg(theil_sen_pct(v, 50)),
        "div": divergenza_cmf(d),
        "squeeze": squeeze_attivo(df),
        "poc": volume_poc(df),
        "atr14": float(atr14.iloc[-1]),
        "ret20": ret20, "rs20": (ret20 - spy_ret20) * 100,
    }


def calcola_regime(spy: pd.DataFrame, vix: Optional[pd.DataFrame]) -> dict:
    c = spy["Close"]
    sma50, sma200 = float(c.rolling(50).mean().iloc[-1]), float(c.rolling(200).mean().iloc[-1])
    last = float(c.iloc[-1])
    vix_last = float(vix["Close"].dropna().iloc[-1]) if vix is not None and not vix.empty else None
    sopra50, sopra200 = last > sma50, last > sma200
    alto_vix = vix_last is not None and vix_last > 25
    if not sopra200 and (not sopra50 or alto_vix):
        etichetta = "RISK-OFF 🔴"
    elif not sopra50 or not sopra200 or alto_vix:
        etichetta = "CAUTELA 🟡"
    else:
        etichetta = "RISK-ON 🟢"
    return {"etichetta": etichetta, "spy_vs_sma200": (last / sma200 - 1) * 100,
            "spy_vs_sma50": (last / sma50 - 1) * 100, "vix": vix_last,
            "risk_off": etichetta.startswith("RISK-OFF")}


# =====================================================================
# 4. FONDAMENTALI: CACHE + RETRY + CIRCUIT BREAKER
# =====================================================================
def _cache_path(tk: str) -> Path:
    return CACHE_DIR / f"{re.sub(r'[^A-Z0-9_-]', '_', tk)}.json"


def _leggi_cache(tk: str) -> tuple[Optional[dict], bool]:
    p = _cache_path(tk)
    if not p.exists():
        return None, False
    try:
        blob = json.loads(p.read_text(encoding="utf-8"))
        fresco = (time.time() - blob["_ts"]) < CFG.fund_ttl_days * 86400
        return blob["data"], fresco
    except Exception:  # noqa: BLE001 - cache corrotta: la ignoriamo
        return None, False


def _scrivi_cache(tk: str, data: dict) -> None:
    try:
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        _cache_path(tk).write_text(json.dumps({"_ts": time.time(), "data": data}), encoding="utf-8")
    except OSError as exc:
        log.debug("cache write %s: %s", tk, exc)


def _scarica_info(tk: str) -> dict:
    info = yf.Ticker(tk).info or {}
    if len(info) < 10:                                   # risposta "vuota" tipica dei rate limit
        raise ValueError("info vuoto/incompleto")
    return {k: info.get(k) for k in FUND_KEYS}


def _scarica_rd_ratio(tk: str, revenue: Optional[float]) -> Optional[float]:
    """R&D / ricavi. Chiamata costosa: solo se il titolo mostra segnali di stress."""
    try:
        fin = yf.Ticker(tk).financials
        if fin is None or fin.empty or not revenue or revenue <= 0:
            return None
        for nome in fin.index:
            if "research" in str(nome).lower():
                serie = fin.loc[nome].sort_index(ascending=False).dropna()
                if not serie.empty:
                    return abs(float(serie.iloc[0])) / revenue
    except Exception:  # noqa: BLE001
        pass
    return None


def scarica_fondamentali(tickers: list[str], stats: Counter) -> dict[str, Optional[dict]]:
    out: dict[str, Optional[dict]] = {}
    lock = threading.Lock()
    stato = {"consec": 0, "abort": False}

    def lavora(tk: str):
        cached, fresco = _leggi_cache(tk)
        if fresco:
            stats["fond_da_cache"] += 1
            return tk, cached
        if stato["abort"]:
            stats["fond_saltati_abort"] += 1
            return tk, cached                              # eventuale cache scaduta meglio di niente
        time.sleep(random.uniform(0.1, 0.4))
        try:
            d = retry(lambda: _scarica_info(tk), tries=3, base=2)
            stress = ((num(d.get("revenueGrowth")) or 0) < -0.30 or (num(d.get("earningsGrowth")) or 0) < -0.35
                      or (num(d.get("debtToEquity")) or 0) > 400)
            d["rd_ratio"] = _scarica_rd_ratio(tk, num(d.get("totalRevenue"))) if stress else None
            _scrivi_cache(tk, d)
            with lock:
                stato["consec"] = 0
                stats["fond_scaricati"] += 1
            return tk, d
        except Exception as exc:  # noqa: BLE001
            with lock:
                stato["consec"] += 1
                stats["fond_errori"] += 1
                if stato["consec"] >= CFG.fund_max_consecutive_fail and not stato["abort"]:
                    stato["abort"] = True
                    stats["fond_abort"] = 1
                    log.error("Circuit breaker: %d errori consecutivi su Yahoo, stop download fondamentali.",
                              stato["consec"])
            log.debug("info %s fallita: %s", tk, exc)
            return tk, cached                              # fallback su cache scaduta (se esiste)

    with ThreadPoolExecutor(max_workers=CFG.info_workers) as ex:
        for tk, d in ex.map(lavora, tickers):
            out[tk] = d
    return out


# =====================================================================
# 5. FAIR VALUE MULTI-MODELLO, QUALITA', VERDETTO
# =====================================================================
def _mediana_pesata(coppie: list[tuple[float, float]]) -> float:
    coppie = sorted(coppie)
    tot, acc = sum(w for _, w in coppie), 0.0
    for v, w in coppie:
        acc += w
        if acc >= tot / 2:
            return v
    return coppie[-1][0]


def stima_fair_value(f: dict, price: float) -> dict:
    """
    Ensemble di modelli indipendenti, combinati con mediana pesata.
    Nessun 'fallback inventato': se non ci sono modelli, FV = None (e lo score lo tratta come neutro).
    """
    sector = f.get("sector")
    industry = (f.get("industry") or "").lower()
    eps_f, eps_t = num(f.get("forwardEps")), num(f.get("trailingEps"))
    modelli: list[tuple[str, float, float]] = []      # (nome, valore, peso)

    # crescita utili attesa: forward vs trailing (piu' stabile del YoY trimestrale di Yahoo)
    g = None
    if eps_f and eps_t and eps_f > 0 and eps_t > 0:
        g = (eps_f / eps_t - 1) * 100
    elif num(f.get("earningsGrowth")) is not None:
        g = num(f.get("earningsGrowth")) * 100

    if sector != "Real Estate" and eps_f and eps_f > 0:
        peso = 0.6 if sector in SETTORI_CICLICI else 1.0          # utili ciclici = stima meno affidabile
        if sector in SETTORI_GROWTH and g is not None and g > 0:
            modelli.append(("P/E~crescita (Lynch)", eps_f * clamp(g, 12, 32), peso))
        else:
            pe = PE_BASE.get(sector, 15)
            modelli.append((f"P/E settoriale {pe}", eps_f * pe, peso))

    shares = num(f.get("sharesOutstanding")) or ((num(f.get("marketCap")) or 0) / price if price else 0)
    fcf = num(f.get("freeCashflow"))
    if sector not in ("Financial Services", "Real Estate") and fcf and fcf > 0 and shares and shares > 0:
        modelli.append((f"FCF yield {CFG.fcf_required_yield:.0%}", fcf / shares / CFG.fcf_required_yield, 0.8))

    bv, roe = num(f.get("bookValue")), num(f.get("returnOnEquity"))
    if sector == "Financial Services" and ("bank" in industry or "insurance" in industry) and bv and bv > 0 and roe is not None:
        pb = clamp((roe - 0.03) / 0.07, 0.6, 2.2)                # P/B giustificato: (ROE-g)/(Ke-g), Ke=10%, g=3%
        modelli.append(("P/B giustificato da ROE", bv * pb, 1.0))

    tgt, n_an = num(f.get("targetMeanPrice")), num(f.get("numberOfAnalystOpinions")) or 0
    if tgt and tgt > 0 and n_an >= CFG.min_analysts:
        modelli.append((f"Target analisti (n={int(n_an)})", tgt, 0.5))

    validi = [(n, v, w) for n, v, w in modelli if CFG.fv_min_ratio * price <= v <= CFG.fv_max_ratio * price]
    if not validi:
        return {"fv": None, "sconto": None, "metodo": "N/D", "conf": "N/D", "n": 0}

    fv = _mediana_pesata([(v, w) for _, v, w in validi])
    valori = [v for _, v, _ in validi]
    disp = (max(valori) - min(valori)) / fv if fv else 9
    if len(validi) >= 3 and disp < 0.4:
        conf = "Alta"
    elif len(validi) >= 2 and disp < 0.8:
        conf = "Media"
    else:
        conf = "Bassa"
    return {"fv": round(fv, 2), "sconto": round((fv - price) / fv * 100, 2),
            "metodo": " + ".join(n for n, _, _ in validi), "conf": conf, "n": len(validi)}


DQ_KEYS = ["forwardEps", "earningsGrowth", "revenueGrowth", "returnOnEquity", "operatingMargins",
           "freeCashflow", "targetMeanPrice", "sector"]


def valuta_qualita(f: Optional[dict], cmf: float) -> dict:
    """
    Filtro anti-dissesto + indice di qualita' (0..1).
    - dati mancanti NON significano 'azienda solida': abbassano Data Quality
    - l'esenzione R&D vale solo se R&D/ricavi > 15% E la cassa copre il burn (o l'FCF e' positivo)
    """
    if not f:
        return {"ok": True, "nota": "Fondamentali non disponibili: valutazione solo tecnica.",
                "qual": None, "dq": "Bassa", "mancanti": len(DQ_KEYS)}

    mancanti = sum(1 for k in DQ_KEYS if f.get(k) in (None, ""))
    dq = "Alta" if mancanti <= 2 else "Media" if mancanti <= 4 else "Bassa"

    settore = f.get("sector")
    finanziario = settore in ("Financial Services", "Real Estate")
    revg, eg = num(f.get("revenueGrowth")), num(f.get("earningsGrowth"))
    de, fcf = num(f.get("debtToEquity")), num(f.get("freeCashflow"))
    cash, debt, ebitda = num(f.get("totalCash")), num(f.get("totalDebt")), num(f.get("ebitda"))
    net_debt = (debt - cash) if (debt is not None and cash is not None) else None
    nde = net_debt / ebitda if (net_debt is not None and ebitda and ebitda > 0) else None
    if de is not None and de < 0:
        de = None                                        # equity negativa (buyback): D/E non interpretabile

    lev_bad = (not finanziario) and (
        (de is not None and de > 400)
        or (nde is not None and nde > 6)
        or (net_debt is not None and net_debt > 0 and ebitda is not None and ebitda <= 0 and fcf is not None and fcf < 0)
    )
    rd = num(f.get("rd_ratio"))
    runway_ok = fcf is not None and (fcf >= 0 or (cash is not None and cash > 1.5 * abs(fcf)))
    esente = rd is not None and rd > 0.15 and runway_ok

    problemi = []
    if revg is not None and revg < -0.30 and cmf < 0.10:
        problemi.append(("Crollo ricavi", "Crollo dei ricavi senza supporto dei flussi."))
    if lev_bad and cmf < 0.05:
        problemi.append(("Leva elevata", "Indebitamento critico senza flussi a favore."))
    if eg is not None and eg < -0.35 and cmf < 0.05:
        problemi.append(("Utili in calo", "Utili in forte calo senza supporto dei flussi."))

    if problemi and not esente:
        return {"ok": False, "nota": problemi[0][1], "qual": None, "dq": dq, "mancanti": mancanti}

    checks = []
    roe, opm, cr = num(f.get("returnOnEquity")), num(f.get("operatingMargins")), num(f.get("currentRatio"))
    if roe is not None:
        checks.append(roe >= 0.10)
    if opm is not None:
        checks.append(opm >= 0.10)
    if fcf is not None:
        checks.append(fcf > 0)
    if nde is not None:
        checks.append(nde < 3)
    elif de is not None and not finanziario:
        checks.append(de < 150)
    if cr is not None and not finanziario:
        checks.append(cr >= 1)
    qual = sum(checks) / len(checks) if len(checks) >= 3 else None

    if problemi and esente:
        nota = f"Segnali di stress ({problemi[0][0].lower()}) giustificati da R&D {rd:.0%} dei ricavi con cassa/FCF adeguati."
    elif qual is None:
        nota = "Dati insufficienti per misurare la qualità (nessun segnale di dissesto)."
    else:
        nota = f"Nessun segnale di dissesto. Indice qualità {qual:.0%} ({sum(checks)}/{len(checks)} test superati)."
    return {"ok": True, "nota": nota, "qual": qual, "dq": dq, "mancanti": mancanti}


def giorni_a_earnings(f: Optional[dict], ora_utc: datetime) -> Optional[int]:
    if not f:
        return None
    for k in ("earningsTimestampStart", "earningsTimestamp"):
        ts = num(f.get(k))
        if ts:
            try:
                d = (datetime.fromtimestamp(ts, tz=timezone.utc) - ora_utc).days
            except (OverflowError, OSError, ValueError):
                continue
            if 0 <= d <= 120:
                return d
    return None


def determina_verdetto(t: dict, short_pct: Optional[float]) -> tuple[Cat, str]:
    cmf = t["cmf"]
    vol_debole = t["trend5"] == Tr5.RAFF or t["reg20"] == VolReg.ESAUR
    if t["vsa"] == Vsa.DISTRIB or (t["div"] == Div.RIBASS and cmf < 0):
        return Cat.EVITARE, "🔴 [EVITARE] Flussi in distribuzione: la pressione di vendita istituzionale prevale."
    if short_pct is not None and short_pct > 10 and cmf > 0.10 and t["squeeze"]:
        if vol_debole or not (t["rvol5"] >= 1.1):
            return Cat.ATTESA, "⚠️ [ATTESA] Squeeze di volatilità con short elevato, ma i volumi non confermano ancora."
        return Cat.BREVE_SQUEEZE, "🔥 [BREVE] Squeeze attivo, short interest alto e volumi in aumento (profilo speculativo)."
    if t["div"] == Div.RIALZ:
        if t["reg50"] == VolReg.ESAUR:
            return Cat.ATTESA, "🔍 [ATTESA] Divergenza rialzista sul CMF, ma i volumi a 50 giorni sono in esaurimento."
        return Cat.BREVE_MEDIO, "🚀 [BREVE/MEDIO] Divergenza rialzista prezzo/CMF con volumi non in esaurimento."
    if t["dist_min60"] <= 3.0 and cmf >= 0:
        return Cat.ACCUMULO, "💎 [ACCUMULO] Prezzo vicino ai minimi a 60 giorni con flussi non negativi."
    if not t["bear"] and cmf > 0.05:
        return Cat.CASSETTO, "🛡️ [CASSETTO] Trend di fondo rialzista (sopra SMA200) con flussi in accumulo."
    if t["bear"] and cmf > 0.05:
        return Cat.PAC, "💎 [PAC] Sotto SMA200 ma con flussi in accumulo: candidato a ingresso graduale."
    return Cat.ATTESA, f"🔍 [ATTESA] Struttura incerta ({t['vsa'].value})."


def valuta_idoneita(f: dict, t: dict, fvr: dict, q: dict, short_pct: Optional[float]) -> str:
    if q["dq"] == "Bassa":
        return "🟡 NEUTRO: dati fondamentali insufficienti, trattare come trade breve."
    fwd_pe, peg = num(f.get("forwardPE")), num(f.get("pegRatio"))
    if short_pct is not None and short_pct > 12 and (fwd_pe is None or fwd_pe > 40):
        return "🔴 SOLO TRADE BREVE: profilo speculativo (short elevato, valutazione tesa)."
    pt = 0
    pt += q["qual"] is not None and q["qual"] >= 0.6
    pt += fwd_pe is not None and 0 < fwd_pe < 25
    pt += peg is not None and 0 < peg < 1.5
    pt += fvr["sconto"] is not None and fvr["sconto"] >= 15 and fvr["conf"] in ("Alta", "Media")
    pt += t["cmf"] > 0.05 or t["vsa"] == Vsa.ACCUMULO
    if pt >= 4:
        return "🟢 IDONEO: qualità, valutazione e flussi coerenti. Candidato per il portafoglio di lungo."
    if pt == 3:
        return "🟡 CON RISERVA: quadro buono ma non completo. Verifica i bilanci prima di un PAC."
    return "🔴 SOLO TRADE BREVE: dinamica tecnica interessante, ma qualità/valutazione non sufficienti per il lungo."


# =====================================================================
# 6. COSTRUZIONE RECORD, SCORE, TOP N
# =====================================================================
def costruisci_record(tk: str, t: dict, f: Optional[dict], regime: dict, ora_utc: datetime):
    q = valuta_qualita(f, t["cmf"])
    if not q["ok"]:
        return None, q["nota"]
    fe = f or {}
    price = t["price"]
    fvr = stima_fair_value(fe, price) if f else {"fv": None, "sconto": None, "metodo": "N/D", "conf": "N/D", "n": 0}
    sp = num(fe.get("shortPercentOfFloat"))
    short_pct = sp * 100 if sp is not None else None
    earn = giorni_a_earnings(f, ora_utc)
    cat, testo = determina_verdetto(t, short_pct)
    if earn is not None and earn <= CFG.earnings_warn_days:
        testo += f" ⚠️ Earnings tra {earn}g."
    if q["dq"] == "Bassa":
        testo += " (dati incompleti)"

    stop = t["support60"] - t["atr14"]
    rischio = price - stop
    rr = (fvr["fv"] - price) / rischio if (fvr["fv"] and fvr["fv"] > price and rischio > 0) else None
    tgt = num(fe.get("targetMeanPrice"))

    pen = 0.0
    pen += 15 if t["vsa"] == Vsa.DISTRIB else 0
    pen += 10 if t["div"] == Div.RIBASS else 0
    pen += 10 if (earn is not None and earn <= CFG.earnings_warn_days) else 0
    pen += 10 if q["dq"] == "Bassa" else 0
    pen += 5 if (regime["risk_off"] and t["bear"]) else 0

    nome = fe.get("shortName") or fe.get("longName") or ""
    r = {
        "TOP 10 OCCASIONI": "-",
        "Rank": None,
        "Score (0-100)": None,
        "Ticker": f"{tk} - {nome}" if nome else tk,
        "Settore": SETTORI_IT.get(fe.get("sector"), fe.get("sector") or "N/D"),
        "Sottosettore": fe.get("industry") or "N/D",
        "VERDETTO DEL CONSULENTE (AZIONE RAPIDA)": testo,
        "Idoneità Portafoglio Personale (PAC/Lungo)": valuta_idoneita(fe, t, fvr, q, short_pct) if f else
        "🟡 NEUTRO: fondamentali non disponibili.",
        "Orizzonte Strategico": ORIZZONTE[cat],
        "Prezzo Attuale ($)": round(price, 2),
        "Fair Value Stimato ($)": fvr["fv"],
        "Sconto su Fair Value (%)": fvr["sconto"],
        "Confidenza FV": fvr["conf"],
        "Metodo Valutazione FV": fvr["metodo"],
        "Stop Suggerito ($)": round(stop, 2) if rischio > 0 else None,
        "Rischio/Rendimento (verso FV)": round(rr, 2) if rr else None,
        "Supporto 60G ($)": round(t["support60"], 2),
        "Distanza dal Minimo 60G (%)": round(t["dist_min60"], 2),
        "Resistenza Max 52W ($)": round(t["max52"], 2),
        "Storno dai Max 52W (%)": round(t["storno"], 1),
        "Distanza da SMA 50 (%)": round(t["dist_sma50"], 2),
        "Distanza da SMA 200 (%)": round(t["dist_sma200"], 2),
        "Trend di Fondo": "🔴 BEAR (Sotto SMA200)" if t["bear"] else "🟢 BULL (Sopra SMA200)",
        "Forza Relativa 20G vs SPY (%)": round(t["rs20"], 2),
        "Target Price Medio ($)": round(tgt, 2) if tgt else None,
        "Upside Atteso (%)": round((tgt / price - 1) * 100, 2) if tgt else None,
        "N. Analisti": int(num(fe.get("numberOfAnalystOpinions"))) if num(fe.get("numberOfAnalystOpinions")) else None,
        "RSI (14, Wilder)": round(t["rsi"], 1),
        "Chaikin Money Flow (CMF)": round(t["cmf"], 3),
        "Divergenza CMF": t["div"].value,
        "Volume Up/Down 20G (x)": round(t["updown"], 2),
        "Volumi 1W (% vs mediana 60G)": round(t["rvol5"] * 100, 1) if math.isfinite(t["rvol5"]) else None,
        "Trend Volumi 5G": t["trend5"].value,
        "Regressione Volumi 20G (1M)": t["reg20"].value,
        "Regressione Volumi 50G (3M)": t["reg50"].value,
        "OBV Trend": "Rialzista (Accumulo)" if t["obv_up"] else "Ribassista (Distribuzione)",
        "Analisi VSA (Flussi)": t["vsa"].value,
        "Squeeze Volatilità": "Attivo 🔥" if t["squeeze"] else "No",
        "Point of Control (POC 60G)": round(t["poc"], 2),
        "Short Interest (%)": round(short_pct, 2) if short_pct is not None else None,
        "Forward P/E": round(num(fe.get("forwardPE")), 2) if num(fe.get("forwardPE")) else None,
        "PEG Ratio": round(num(fe.get("pegRatio")), 2) if num(fe.get("pegRatio")) else None,
        "Earnings tra (gg)": earn,
        "Data Quality": q["dq"],
        "Analisi Fondamentale / Nota": q["nota"],
        # --- campi interni (non esportati in Excel)
        "_symbol": tk, "_cat": cat, "_azionabile": cat in AZIONABILI, "_sconto": fvr["sconto"],
        "_conf_w": {"Alta": 1.0, "Media": 0.8, "Bassa": 0.5}.get(fvr["conf"], 1.0),
        "_cmf": t["cmf"], "_updown": t["updown"], "_qual": q["qual"], "_rs20": t["rs20"],
        "_div_bull": float(t["div"] == Div.RIALZ), "_rsi_rimb": float(t["rsi_rimbalzo"]),
        "_near_min": float(t["dist_min60"] <= 3.0), "_pen": pen, "_sector_en": fe.get("sector") or "N/D",
    }
    return r, None


def applica_score(df: pd.DataFrame) -> pd.DataFrame:
    """
    Score = media pesata di PERCENTILI cross-sectional (nessuna soglia assoluta, nessuna saturazione a 100).
    Una sola famiglia per fattore per evitare doppio conteggio: CMF e Up/Down volume per i flussi.
    OBV/VSA/regressioni volumi restano nel report ma non sono sommati allo score.
    """
    for c in ("_sconto", "_cmf", "_updown", "_qual", "_rs20", "_conf_w", "_pen"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    pct = lambda col: df[col].rank(pct=True)  # noqa: E731

    value = np.where(df["_sconto"].notna(), pct("_sconto").fillna(0.25) * df["_conf_w"], 0.25)
    flows = 0.5 * pct("_cmf").fillna(0.5) + 0.5 * pct("_updown").fillna(0.5)
    quality = pct("_qual").fillna(0.4)
    setup = (0.35 * pct("_rs20").fillna(0.5) + 0.25 * df["_div_bull"] + 0.20 * df["_rsi_rimb"]
             + 0.20 * df["_near_min"])
    score = 100 * (0.30 * value + 0.30 * flows + 0.20 * quality + 0.20 * setup) - df["_pen"]
    df["Score (0-100)"] = score.clip(0, 100).round(1)
    return df


def seleziona_top(df: pd.DataFrame) -> pd.DataFrame:
    df = df.sort_values(["Score (0-100)", "_sconto"], ascending=[False, False], na_position="last").reset_index(drop=True)
    scelti, per_settore = [], Counter()
    for idx, row in df.iterrows():
        if CFG.only_actionable_top and not row["_azionabile"]:
            continue
        if per_settore[row["_sector_en"]] >= CFG.max_per_sector:
            continue
        scelti.append(idx)
        per_settore[row["_sector_en"]] += 1
        if len(scelti) >= CFG.top_n:
            break
    for pos, idx in enumerate(scelti, 1):
        df.at[idx, "TOP 10 OCCASIONI"] = f"⭐ TOP {pos}"
    ordine = {idx: pos for pos, idx in enumerate(scelti)}
    df["_ord"] = df.index.map(lambda i: ordine.get(i, 10_000 + i))
    df = df.sort_values("_ord").reset_index(drop=True)
    df["Rank"] = np.arange(1, len(df) + 1)
    return df.drop(columns="_ord")


# =====================================================================
# 7. STORICO (per forward test) E CONFRONTO CON IERI
# =====================================================================
HIST_COLS = ["date", "symbol", "rank", "top10", "score", "categoria", "azionabile", "price",
             "fv", "sconto", "conf", "sector", "cmf", "storno", "regime"]


def salva_storico(df: pd.DataFrame, data_mercato: str, regime: dict) -> None:
    h = pd.DataFrame({
        "date": data_mercato, "symbol": df["_symbol"], "rank": df["Rank"], "top10": df["TOP 10 OCCASIONI"] != "-",
        "score": df["Score (0-100)"], "categoria": df["_cat"].map(lambda c: c.value), "azionabile": df["_azionabile"],
        "price": df["Prezzo Attuale ($)"], "fv": df["Fair Value Stimato ($)"], "sconto": df["Sconto su Fair Value (%)"],
        "conf": df["Confidenza FV"], "sector": df["_sector_en"], "cmf": df["Chaikin Money Flow (CMF)"],
        "storno": df["Storno dai Max 52W (%)"], "regime": regime["etichetta"].split()[0],
    })[HIST_COLS]
    HISTORY_FILE.parent.mkdir(parents=True, exist_ok=True)
    if HISTORY_FILE.exists():
        vecchio = pd.read_csv(HISTORY_FILE, dtype={"date": str})
        vecchio = vecchio[vecchio["date"] != data_mercato]          # idempotente: rilancio nello stesso giorno = sostituisce
        h = pd.concat([vecchio, h], ignore_index=True)
    h.to_csv(HISTORY_FILE, index=False)


def top_precedente(data_mercato: str) -> Optional[set]:
    if not HISTORY_FILE.exists():
        return None
    try:
        h = pd.read_csv(HISTORY_FILE, dtype={"date": str})
        h = h[h["date"] < data_mercato]
        if h.empty:
            return None
        ultimo = h[h["date"] == h["date"].max()]
        return set(ultimo.loc[ultimo["top10"].astype(str).str.lower() == "true", "symbol"])
    except Exception:  # noqa: BLE001
        return None


# =====================================================================
# 8. EXCEL
# =====================================================================
FONT = "Arial"
LEGENDA = [
    ("Score (0-100)", "Media pesata di percentili sull'universo del giorno: valore 30%, flussi 30%, qualità 20%, setup/forza relativa 20%, meno penalità (distribuzione, earnings vicini, dati incompleti, risk-off). E' un ranking relativo, non una probabilità."),
    ("TOP 10", "Migliori per score tra i soli verdetti azionabili (no ATTESA/EVITARE), max 3 per settore."),
    ("Fair Value / Confidenza FV", "Mediana pesata di più modelli (P/E~crescita o settoriale, FCF yield, P/B da ROE per banche/assicurazioni, target analisti se n>=5). Confidenza Alta = >=3 modelli concordi; Bassa = modelli discordi o uno solo. Sconto = (FV-Prezzo)/FV."),
    ("Stop Suggerito / Rischio-Rendimento", "Stop = minimo 60G - 1 ATR(14). R/R = (FV - prezzo) / (prezzo - stop). Indicativo, non sostituisce il tuo piano di rischio."),
    ("CMF", "Chaikin Money Flow 20G: >0 pressione d'acquisto, <0 vendita."),
    ("Volume Up/Down 20G", "Volume nei giorni di rialzo / volume nei giorni di ribasso (ultimi 20). >1 = accumulo direzionale."),
    ("Regressione Volumi", "Pendenza robusta (Theil-Sen) del log-volume su 20/50 giorni; soglia ±0,5%/giorno."),
    ("Divergenza CMF", "Confronto tra gli ultimi due pivot di prezzo e di CMF (finestra 45 barre)."),
    ("Forza Relativa 20G vs SPY", "Rendimento a 20 giorni del titolo meno quello di SPY (punti %)."),
    ("Data Quality", "Alta/Media/Bassa in base ai campi Yahoo disponibili. Bassa = penalità sullo score e verdetto marcato 'dati incompleti'."),
    ("OBV, POC, VSA, Squeeze", "Mostrati per contesto; non sommati allo score per evitare doppio conteggio."),
    ("Limiti", "Dati Yahoo Finance gratuiti (non point-in-time, possibili errori/ritardi). Strumento di screening a scopo informativo, non consulenza finanziaria."),
]


def _xl(v):
    if v is None:
        return None
    if isinstance(v, Enum):
        return v.value
    if isinstance(v, (np.bool_,)):
        return bool(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        return None if not math.isfinite(v) else float(v)
    return v


def genera_excel(df: pd.DataFrame, info_run: list[tuple[str, str]], path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consulente Smart Money"
    cols = [c for c in df.columns if not c.startswith("_")]
    ws.append(cols)
    for row in df[cols].itertuples(index=False):
        ws.append([_xl(v) for v in row])

    header_fill = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
    top_fill = PatternFill("solid", start_color="D9EAD3", end_color="D9EAD3")
    side = Side(style="thin", color="D9D9D9")
    border = Border(left=side, right=side, top=side, bottom=side)
    testo_lungo = {"VERDETTO DEL CONSULENTE (AZIONE RAPIDA)": 55, "Idoneità Portafoglio Personale (PAC/Lungo)": 50,
                   "Analisi Fondamentale / Nota": 55, "Metodo Valutazione FV": 40, "Ticker": 32,
                   "Sottosettore": 26, "Settore": 26}

    for cell in ws[1]:
        cell.fill, cell.border = header_fill, border
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 45

    for r_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        is_top = str(ws.cell(row=r_idx, column=1).value or "-") != "-"
        for cell in row:
            name = cols[cell.column - 1]
            cell.font = Font(name=FONT, size=10, bold=is_top and name in ("TOP 10 OCCASIONI", "Ticker"))
            cell.border = border
            if is_top:
                cell.fill = top_fill
            if name in testo_lungo:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.000" if "CMF" in name else "0.0" if ("(%)" in name or "Score" in name or "RSI" in name) else "0.00"

    for i, name in enumerate(cols, start=1):
        largh = testo_lungo.get(name) or min(max(len(name) * 0.8, 12), 20)
        ws.column_dimensions[get_column_letter(i)].width = largh

    n = ws.max_row
    for nome_col, inverso in (("Score (0-100)", False), ("Chaikin Money Flow (CMF)", False), ("Sconto su Fair Value (%)", False)):
        if nome_col in cols and n > 2:
            L = get_column_letter(cols.index(nome_col) + 1)
            ws.conditional_formatting.add(
                f"{L}2:{L}{n}",
                ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
                               mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    ws.freeze_panes = ws.cell(row=2, column=cols.index("Ticker") + 2)
    ws.auto_filter.ref = ws.dimensions

    for titolo, righe, w in (("Legenda", LEGENDA, (32, 120)), ("Info Run", info_run, (32, 80))):
        s = wb.create_sheet(titolo)
        s.append(["Voce", "Dettaglio"])
        for a, b in righe:
            s.append([a, b])
        for c in s[1]:
            c.fill, c.font = header_fill, Font(name=FONT, bold=True, color="FFFFFF")
        for row in s.iter_rows(min_row=2):
            row[0].font = Font(name=FONT, bold=True, size=10)
            row[1].font = Font(name=FONT, size=10)
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
            row[0].alignment = Alignment(vertical="top")
        s.column_dimensions["A"].width, s.column_dimensions["B"].width = w

    wb.save(path)
    return path


# =====================================================================
# 9. EMAIL
# =====================================================================
def _html_report(df, regime, precedente, info_righe, parziale, data_mercato) -> tuple[str, str]:
    top = df[df["TOP 10 OCCASIONI"] != "-"]
    e = html.escape
    righe_html, righe_txt = [], []
    for _, r in top.iterrows():
        fv = f"{r['Sconto su Fair Value (%)']:.0f}% ({r['Confidenza FV']})" if pd.notna(r["Sconto su Fair Value (%)"]) else "n/d"
        stop = f"{r['Stop Suggerito ($)']:.2f}" if pd.notna(r["Stop Suggerito ($)"]) else "n/d"
        righe_html.append(
            f"<tr><td>{e(str(r['TOP 10 OCCASIONI']))}</td><td><b>{e(str(r['Ticker']))}</b></td>"
            f"<td align='center'>{r['Score (0-100)']:.0f}</td><td>{e(str(r['VERDETTO DEL CONSULENTE (AZIONE RAPIDA)']))}</td>"
            f"<td align='right'>{r['Prezzo Attuale ($)']:.2f}</td><td align='center'>{fv}</td><td align='right'>{stop}</td></tr>")
        righe_txt.append(f"{r['TOP 10 OCCASIONI']} {r['Ticker']} | score {r['Score (0-100)']:.0f} | {r['VERDETTO DEL CONSULENTE (AZIONE RAPIDA)']}")

    oggi = set(top["_symbol"])
    novita = ""
    if precedente is not None:
        nuovi, usciti = sorted(oggi - precedente), sorted(precedente - oggi)
        novita = (f"<p><b>Novità vs seduta precedente</b> — entrati: {e(', '.join(nuovi) or 'nessuno')}; "
                  f"usciti: {e(', '.join(usciti) or 'nessuno')}.</p>")
    avviso = ("<p style='color:#b45f06'><b>⚠️ Run parziale:</b> Yahoo ha risposto male a una parte delle richieste; "
              "i risultati potrebbero essere incompleti.</p>") if parziale else ""
    vix = f", VIX {regime['vix']:.1f}" if regime["vix"] else ""
    corpo = f"""<html><body style="font-family:Arial,sans-serif;font-size:13px">
<h2>🧠 Money Flow Pro — seduta del {e(data_mercato)}</h2>
<p><b>Regime di mercato:</b> {e(regime['etichetta'])} (SPY {regime['spy_vs_sma200']:+.1f}% vs SMA200{vix}).
Candidati analizzati: <b>{len(df)}</b>, azionabili: <b>{int(df['_azionabile'].sum())}</b>.</p>
{avviso}{novita}
<table border="1" cellpadding="5" cellspacing="0" style="border-collapse:collapse;border-color:#ccc">
<tr style="background:#1F4E79;color:#fff"><th>#</th><th>Ticker</th><th>Score</th><th>Verdetto</th><th>Prezzo $</th><th>Sconto FV</th><th>Stop $</th></tr>
{''.join(righe_html) or "<tr><td colspan='7'>Nessun candidato azionabile oggi.</td></tr>"}
</table>
<p style="color:#666;font-size:11px">Report completo nell'allegato Excel. Strumento di screening a scopo informativo:
dati Yahoo Finance gratuiti, non è consulenza finanziaria.<br>{e(' | '.join(f'{a}: {b}' for a, b in info_righe[:6]))}</p>
</body></html>"""
    testo = f"Money Flow Pro {data_mercato} - regime {regime['etichetta']}\n" + "\n".join(righe_txt)
    return corpo, testo


def invia_email(oggetto: str, html_body: str, testo: str, allegato: Optional[Path] = None) -> bool:
    if not (SENDER_EMAIL and APP_PASSWORD and RECEIVER_EMAIL):
        log.warning("Credenziali email non configurate: nessun invio.")
        return True
    msg = MIMEMultipart("mixed")
    msg["From"], msg["To"], msg["Subject"] = SENDER_EMAIL, RECEIVER_EMAIL, oggetto
    alt = MIMEMultipart("alternative")
    alt.attach(MIMEText(testo, "plain", "utf-8"))
    alt.attach(MIMEText(html_body, "html", "utf-8"))
    msg.attach(alt)
    if allegato and allegato.exists():
        part = MIMEApplication(allegato.read_bytes(), Name=allegato.name)
        part["Content-Disposition"] = f'attachment; filename="{allegato.name}"'
        msg.attach(part)
    ctx = ssl.create_default_context()
    for tentativo in range(3):
        try:
            with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ctx, timeout=30) as s:
                s.login(SENDER_EMAIL, APP_PASSWORD)
                s.send_message(msg)
            log.info("✅ Email inviata.")
            return True
        except Exception as exc:  # noqa: BLE001
            log.warning("Invio email fallito (%d/3): %s", tentativo + 1, type(exc).__name__)
            time.sleep(5 * (tentativo + 1))
    return False


# =====================================================================
# 10. ORCHESTRAZIONE
# =====================================================================
def esegui(limit: Optional[int], no_email: bool) -> int:
    t0 = time.time()
    stats: Counter = Counter()
    ora_utc = datetime.now(timezone.utc)
    ora_ny = datetime.now(NY)

    tickers = ottieni_universo()
    if limit:
        tickers = tickers[:limit]
    stats["universo"] = len(tickers)
    log.info("Universo: %d ticker", len(tickers))

    bench = scarica_prezzi(["SPY", "^VIX"], stats)
    spy = pulisci_ohlcv(bench.get("SPY"), ora_ny)
    if spy is None or len(spy) < 200:
        raise RuntimeError("Dati SPY non disponibili: impossibile determinare regime e data di riferimento.")
    vix = pulisci_ohlcv(bench.get("^VIX"), ora_ny) if "^VIX" in bench else None
    regime = calcola_regime(spy, vix)
    data_mercato = spy.index[-1].date().isoformat()
    spy_ret20 = float(spy["Close"].iloc[-1] / spy["Close"].iloc[-21] - 1)
    log.info("Seduta di riferimento %s | regime %s", data_mercato, regime["etichetta"])

    prezzi = scarica_prezzi(tickers, stats)

    # ---- fase 1: validazione + filtro tecnico
    candidati: dict[str, dict] = {}
    limite_stale = spy.index[-1] - timedelta(days=CFG.stale_days)
    for tk, raw in prezzi.items():
        d = pulisci_ohlcv(raw, ora_ny)
        if d is None or len(d) < CFG.min_bars:
            stats["scartati_storia_corta"] += 1
            continue
        if d.index[-1] < limite_stale:
            stats["scartati_stale"] += 1
            continue
        if float(d["Close"].iloc[-1]) < CFG.min_price or float((d["Close"] * d["Volume"]).iloc[-20:].mean()) < CFG.min_dollar_volume:
            stats["scartati_liquidita"] += 1
            continue
        storno = (float(d["High"].max()) - float(d["Close"].iloc[-1])) / float(d["High"].max()) * 100
        if storno < CFG.storno_min:
            stats["scartati_storno_basso"] += 1
            continue
        if storno > CFG.storno_max:
            stats["scartati_storno_eccessivo"] += 1
            continue
        try:
            t = calcola_tecnici(d, spy_ret20)
        except Exception as exc:  # noqa: BLE001
            log.warning("Indicatori %s falliti: %s", tk, exc)
            stats["errori_indicatori"] += 1
            continue
        if t is None:
            stats["scartati_indicatori_nan"] += 1
            continue
        candidati[tk] = t
    log.info("Candidati dopo filtro tecnico: %d", len(candidati))

    # ---- fase 2: fondamentali solo per i candidati
    fond = scarica_fondamentali(list(candidati), stats) if candidati else {}

    # ---- fase 3: qualita', FV, verdetto
    records = []
    for tk, t in candidati.items():
        try:
            rec, motivo = costruisci_record(tk, t, fond.get(tk), regime, ora_utc)
        except Exception as exc:  # noqa: BLE001
            log.warning("Record %s fallito: %s", tk, exc)
            log.debug(traceback.format_exc())
            stats["errori_record"] += 1
            continue
        if rec is None:
            stats["scartati_bilancio"] += 1
            log.debug("%s scartato: %s", tk, motivo)
            continue
        records.append(rec)

    parziale = bool(stats["fond_abort"]) or stats["prezzi_mancanti"] > 0.05 * max(1, stats["universo"]) \
        or stats["fond_errori"] > 0.10 * max(1, len(candidati))
    info_run = [("Data seduta", data_mercato), ("Regime", regime["etichetta"]),
                ("SPY vs SMA200 (%)", f"{regime['spy_vs_sma200']:+.2f}"),
                ("VIX", f"{regime['vix']:.1f}" if regime["vix"] else "N/D"),
                ("Run parziale", "SI" if parziale else "no"),
                ("Durata (s)", f"{time.time() - t0:.0f}")] + [(k, str(v)) for k, v in sorted(stats.items())]
    for k, v in info_run:
        log.info("%-28s %s", k, v)

    if not records:
        corpo = (f"<p>Nessun titolo ha superato i filtri nella seduta del {data_mercato} "
                 f"(regime {regime['etichetta']}).</p><p>{html.escape(' | '.join(f'{a}: {b}' for a, b in info_run))}</p>")
        if not no_email:
            invia_email(f"🧠 Money Flow Pro ({data_mercato}) - nessuna occasione", corpo, "Nessuna occasione.")
        return 0

    df = pd.DataFrame(records)
    df = applica_score(df)
    df = seleziona_top(df)
    precedente = top_precedente(data_mercato)
    salva_storico(df, data_mercato, regime)

    xlsx = genera_excel(df, info_run, Path(f"Consulente_Smart_Money_{data_mercato}.xlsx"))
    top = df[df["TOP 10 OCCASIONI"] != "-"]
    log.info("Top %d: %s", len(top), ", ".join(top["_symbol"]))

    if no_email:
        return 0
    corpo, testo = _html_report(df, regime, precedente, info_run, parziale, data_mercato)
    ok = invia_email(f"🧠 Money Flow Pro ({data_mercato}) - Top {len(top)} su {len(df)} candidati", corpo, testo, xlsx)
    return 0 if ok else 1


def main() -> int:
    ap = argparse.ArgumentParser(description="Money Flow Pro v2")
    ap.add_argument("--limit", type=int, default=None, help="analizza solo i primi N ticker (test)")
    ap.add_argument("--no-email", action="store_true", help="non inviare l'email")
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args()
    logging.basicConfig(level=logging.DEBUG if a.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)-7s %(message)s", datefmt="%H:%M:%S")
    try:
        return esegui(a.limit, a.no_email)
    except Exception as exc:  # noqa: BLE001
        log.error("Errore fatale: %s", exc)
        tb = traceback.format_exc()
        log.error(tb)
        if not a.no_email:
            invia_email("❌ Money Flow Pro - errore", f"<pre>{html.escape(tb[-3000:])}</pre>", tb[-3000:])
        return 2


if __name__ == "__main__":
    sys.exit(main())
